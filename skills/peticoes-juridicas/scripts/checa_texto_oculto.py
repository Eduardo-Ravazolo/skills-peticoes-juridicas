"""Verifica se um .docx tem texto oculto, branco ou minúsculo (risco de 'prompt injection').

Uso: python3 checa_texto_oculto.py PETICAO.docx
"""
import sys
from docx import Document
from docx.shared import RGBColor

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def runs_all(doc):
    for p in doc.paragraphs:
        yield from p.runs
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield from p.runs
    for s in doc.sections:
        for part in (s.header, s.footer):
            for p in part.paragraphs:
                yield from p.runs


def main(path):
    d = Document(path)
    prob = []
    for r in runs_all(d):
        f = r.font
        if not r.text.strip():
            continue
        if f.hidden or r._element.find(".//" + W + "vanish") is not None:
            prob.append(("oculto", r.text))
        if f.color is not None and f.color.type is not None and f.color.rgb == RGBColor(0xFF, 0xFF, 0xFF):
            prob.append(("branco", r.text))
        if f.size is not None and f.size.pt < 8:
            prob.append(("minúsculo", r.text))
    if prob:
        for tipo, txt in prob:
            print(f"[{tipo}] {txt[:120]}")
        sys.exit(1)
    print("OK: sem texto oculto, branco ou minúsculo")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
