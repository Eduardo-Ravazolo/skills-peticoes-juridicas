"""Gera uma petição no modelo .docx do escritório a partir de uma lista de blocos JSON.

Uso:
    python3 gerar_peticao.py config/escritorio.json conteudo.json SAIDA.docx

- Mantém cabeçalho, rodapé, margens e estilos do modelo .docx indicado no config; apaga o corpo e reescreve.
- O modelo precisa ter os estilos "Heading 2" e "Heading 3" (Título 2 / Título 3).

Blocos (lista JSON), cada um {"t": tipo, "x": texto}:
  enderecamento | qualificacao | acao | texto | secao | subsecao | citacao |
  ref (referência da citação, ex. "(TRF4, AC ..., julgado em ...) (grifei)") |
  pedido | subpedido | nota | print | tabela ("x": [[rótulo, valor], ...], "titulo": opcional) |
  imagem ("x": caminho, "largura_cm" e "legenda" opcionais) | fecho ("x": "dd de mês de aaaa") | quebra
Negrito inline: **trecho**.  Itálico inline: _trecho_ (não usar dentro de citação, que já é itálica).
"""
import json, re, sys, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

J, C = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER
CFG = {}


def runs(p, text, bold=False, italic=False, size=None, hl=False):
    size = size or CFG.get("tamanho_fonte", 12)
    for part in re.split(r"(\*\*.+?\*\*|(?<!\w)_.+?_(?!\w))", text):
        if not part:
            continue
        b, i = bold, italic
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            part, i = part[1:-1], True
        r = p.add_run(part)
        r.bold, r.italic = b or None, i or None
        r.font.name, r.font.size = CFG.get("fonte", "Times New Roman"), Pt(size)
        r._element.rPr.rFonts.set(qn("w:hAnsi"), CFG.get("fonte", "Times New Roman"))
        if hl:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def par(doc, text="", align=J, fi="padrao", li=None, bold=False, italic=False, size=None, hl=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    p.alignment = align
    if fi == "padrao":
        fi = CFG.get("recuo_primeira_linha_cm", 1.5)
    if fi is not None:
        pf.first_line_indent = Cm(fi)
    if li is not None:
        pf.left_indent = Cm(li)
    if text:
        runs(p, text, bold, italic, size, hl)
    return p


def blank(doc, n=1):
    for _ in range(n):
        par(doc, "")


def heading(doc, text, level):
    p = doc.add_paragraph(style="Heading %d" % level)
    pf = p.paragraph_format
    pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
    base = CFG.get("tamanho_fonte", 12)
    runs(p, text, bold=True, size=base + 2 if level == 2 else base)
    return p


def set_borders(table):
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + e)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:color"), "000000")
        b.append(el)
    table._tbl.tblPr.append(b)


def tabela(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    set_borders(t)
    for lab, val in rows:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Cm(5.5), Cm(11)
        for cell, txt, b in ((cells[0], lab, True), (cells[1], val, False)):
            p = cell.paragraphs[0]; p.alignment = J
            runs(p, txt, bold=b, size=CFG.get("tamanho_fonte", 12) - 2)
    return t


def main(config, conteudo, saida):
    global CFG
    CFG = json.load(open(config, encoding="utf-8"))
    modelo = CFG["modelo_docx"]
    if not os.path.isabs(modelo):
        modelo = os.path.join(os.path.dirname(os.path.abspath(config)), "..", modelo)
    doc = Document(modelo)
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)
    citacao_li = CFG.get("recuo_citacao_cm", 3.0)
    blocos = json.load(open(conteudo, encoding="utf-8"))
    for k, b in enumerate(blocos):
        nxt = blocos[k + 1]["t"] if k + 1 < len(blocos) else None
        t, x = b["t"], b.get("x", "")
        if t == "enderecamento":
            par(doc, x, fi=None, bold=True); blank(doc, CFG.get("linhas_em_branco_apos_enderecamento", 10))
        elif t == "qualificacao":
            par(doc, x)
        elif t == "acao":
            blank(doc); par(doc, x, align=C, fi=None, bold=True); blank(doc)
        elif t == "texto":
            par(doc, x)
        elif t == "secao":
            blank(doc); heading(doc, x, 2)
        elif t == "subsecao":
            blank(doc); heading(doc, x, 3)
        elif t == "citacao":
            par(doc, x, li=citacao_li, italic=True)
        elif t == "ref":
            par(doc, x, fi=None, li=citacao_li)
        elif t == "pedido":
            par(doc, x, bold=True)
        elif t == "subpedido":
            par(doc, x, li=1.0, fi=1.0, bold=True)
        elif t == "nota":
            par(doc, x, italic=True)
        elif t == "print":
            par(doc, "[INSERIR PRINT: %s]" % x, align=C, fi=None, bold=True, hl=True)
        elif t == "imagem":
            p = par(doc, "", align=C, fi=None)
            p.add_run().add_picture(x, width=Cm(b.get("largura_cm", 14)))
            if b.get("legenda"):
                par(doc, b["legenda"], align=C, fi=None, italic=True, size=CFG.get("tamanho_fonte", 12) - 2)
        elif t == "tabela":
            if b.get("titulo"):
                par(doc, b["titulo"], fi=None, bold=True)
            tabela(doc, x)
        elif t == "fecho":
            blank(doc); par(doc, CFG.get("fecho", "Nestes Termos, Pede Deferimento.")); blank(doc)
            par(doc, "%s, %s." % (CFG.get("cidade_uf", "Cidade/UF"), x)); blank(doc, 2)
            for i, a in enumerate(CFG.get("assinaturas", [])):
                if i:
                    blank(doc)
                par(doc, a["nome"], align=C, fi=None, bold=True)
                if a.get("oab"):
                    par(doc, a["oab"], align=C, fi=None, bold=True)
        elif t == "quebra":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            continue
        if t == "citacao" and nxt == "ref":
            continue
        if t in ("qualificacao", "texto", "citacao", "ref", "pedido", "subpedido", "nota", "print", "imagem", "tabela", "secao", "subsecao"):
            blank(doc)
    doc.save(saida)


if __name__ == "__main__":
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(*sys.argv[1:4])
